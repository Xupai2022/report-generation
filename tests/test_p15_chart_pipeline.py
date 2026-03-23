from pathlib import Path

from openpyxl import Workbook

from mss_ai_ppt_sample_assets.backend.models.inputs import TenantInput
from mss_ai_ppt_sample_assets.backend.modules.excel_handler import ExcelDataExtractor
from mss_ai_ppt_sample_assets.backend.modules.llm_orchestrator import LLMOrchestratorV2
from mss_ai_ppt_sample_assets.backend.modules.ppt_generator import PPTGeneratorV2


class _DummyTemplateRepo:
    pass


class _DummyRun:
    def __init__(self, text: str):
        self.text = text


class _DummyParagraph:
    def __init__(self, text: str):
        self.runs = [_DummyRun(text)]


class _DummyTextFrame:
    def __init__(self, text: str):
        self.paragraphs = [_DummyParagraph(text)]


class _DummyShape:
    def __init__(self, text: str, left: int, top: int):
        self.shape_type = 1
        self.has_text_frame = True
        self.text_frame = _DummyTextFrame(text)
        self.left = left
        self.top = top


class _DummySlide:
    def __init__(self, shapes):
        self.shapes = shapes


def test_extract_data_includes_p15_top5_chart_payloads(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = ExcelDataExtractor.CLASSIC_REQUIRED_SHEET

    region_rows = [("中国", 12), ("美国", 8), ("新加坡", 5), ("德国", 3), ("日本", 1)]
    attack_type_rows = [("扫描探测", 20), ("暴力破解", 16), ("漏洞利用", 9), ("恶意外联", 6), ("木马下载", 2)]
    host_rows = [("10.0.0.1", 11), ("10.0.0.2", 9), ("10.0.0.3", 7), ("10.0.0.4", 4), ("10.0.0.5", 2)]

    for idx, (label, value) in enumerate(region_rows, start=85):
        ws.cell(idx, 6, label)
        ws.cell(idx, 7, value)
    for idx, (label, value) in enumerate(attack_type_rows, start=85):
        ws.cell(idx, 9, label)
        ws.cell(idx, 10, value)
    for idx, (label, value) in enumerate(host_rows, start=85):
        ws.cell(idx, 12, label)
        ws.cell(idx, 13, value)

    path = tmp_path / "p15.xlsx"
    wb.save(path)

    data = ExcelDataExtractor.extract_data(path)
    threat_effectiveness = data["threat_effectiveness"]

    assert threat_effectiveness["attack_source_region_top5"] == {
        "categories": ["中国", "美国", "新加坡", "德国", "日本"],
        "values": [12, 8, 5, 3, 1],
    }
    assert threat_effectiveness["attack_type_top5"]["values"] == [20, 16, 9, 6, 2]
    assert threat_effectiveness["externally_attacked_hosts_top5"]["categories"][0] == "10.0.0.1"


def test_llm_orchestrator_formats_p15_chart_payloads():
    orchestrator = LLMOrchestratorV2.__new__(LLMOrchestratorV2)
    tenant_input = TenantInput(
        raw={
            "threat_effectiveness": {
                "attack_source_region_top5": {
                    "categories": ["中国", "美国"],
                    "values": [12, 8],
                },
                "externally_attacked_hosts_top5": {
                    "categories": ["10.0.0.1", "10.0.0.2"],
                    "values": [11, 9],
                },
            }
        }
    )

    pie_data = orchestrator._extract_chart_data(
        tenant_input,
        {"data_source": "threat_effectiveness.attack_source_region_top5"},
        "P15_pie_1",
    )
    bar_data = orchestrator._extract_chart_data(
        tenant_input,
        {"data_source": "threat_effectiveness.externally_attacked_hosts_top5", "series_name": "攻击次数"},
        "P15_bar",
    )

    assert pie_data == {"categories": ["中国", "美国"], "values": [12, 8]}
    assert bar_data == {
        "categories": ["10.0.0.1", "10.0.0.2"],
        "series": [{"name": "攻击次数", "values": [11, 9]}],
    }


def test_ppt_generator_finds_exact_p15_placeholder_token():
    generator = PPTGeneratorV2(_DummyTemplateRepo())
    slide = _DummySlide(
        [
            _DummyShape("{{P15_pie_1}}", 100, 200),
            _DummyShape("{{{P15_pie_2}}}", 300, 400),
            _DummyShape("{{P15_bar}}", 500, 600),
        ]
    )

    shape, position = generator._find_placeholder_shape(slide, "P15_pie_2")

    assert shape is slide.shapes[1]
    assert position == (300, 400)
