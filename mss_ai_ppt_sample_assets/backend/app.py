from fastapi import FastAPI, HTTPException, WebSocket, WebSocketDisconnect, UploadFile, File
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pathlib import Path
from pydantic import BaseModel, Field
from typing import Optional
import logging
from datetime import datetime
import asyncio
import json
import openpyxl

from mss_ai_ppt_sample_assets.backend.services.report_service import (
    ReportService,
    InputNotFoundError,
    SlideSpecNotFoundError,
)
from mss_ai_ppt_sample_assets.backend.modules.template_loader import TemplateNotFoundError
from mss_ai_ppt_sample_assets.backend.modules.llm_orchestrator import LLMGenerationError
from mss_ai_ppt_sample_assets.backend import config
from mss_ai_ppt_sample_assets.backend.websocket_support import WebSocketManager
from mss_ai_ppt_sample_assets.backend.websocket_support import WebSocketManager
from openai import RateLimitError

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[
        logging.StreamHandler(),  # Console output
    ]
)
logger = logging.getLogger(__name__)

app = FastAPI(title="MSS AI PPT Backend", version="0.1.0")
logger.info("Initializing MSS AI PPT Backend...")
logger.info(f"LLM Enabled: {config.settings.enable_llm}")
logger.info(f"OpenAI Model: {config.settings.openai_model}")
logger.info(f"OpenAI Base URL: {config.settings.openai_base_url}")

# OpenAI Concurrency Limiter
# Limit concurrent LLM requests to prevent API rate limiting
# Adjust MAX_CONCURRENT_LLM_REQUESTS based on your OpenAI plan's rate limits
MAX_CONCURRENT_LLM_REQUESTS = 3  # Allow max 3 concurrent LLM requests
llm_semaphore = asyncio.Semaphore(MAX_CONCURRENT_LLM_REQUESTS)
logger.info(f"LLM Concurrency Limiter: max {MAX_CONCURRENT_LLM_REQUESTS} concurrent requests")

service = ReportService()

# WebSocket Manager for real-time progress updates
ws_manager = WebSocketManager()
logger.info("WebSocket Manager initialized")

# WebSocket Manager for real-time progress updates
ws_manager = WebSocketManager()
logger.info("WebSocket Manager initialized")
app.mount("/static/previews", StaticFiles(directory=config.PREVIEWS_DIR), name="previews")

# 简单的前端静态页面（无需 npm），挂载在 /ui
FRONTEND_DIR = Path(__file__).parent / "frontend"
if FRONTEND_DIR.exists():
    app.mount("/ui", StaticFiles(directory=FRONTEND_DIR, html=True), name="frontend")


class GenerateRequest(BaseModel):
    input_id: str
    template_id: str
    use_mock: bool = False
    session_id: Optional[str] = None  # Optional: will be auto-generated if not provided
    client_id: Optional[str] = None  # Optional: WebSocket client ID
    client_id: Optional[str] = None  # Optional: WebSocket client ID


class RewriteRequest(BaseModel):
    job_id: str
    slide_key: str
    new_content: dict


class PreviewRequest(BaseModel):
    job_id: str
    regenerate_if_missing: bool = True


@app.websocket("/ws/{client_id}")
async def websocket_endpoint(websocket: WebSocket, client_id: str):
    await ws_manager.connect(websocket, client_id)
    try:
        while True:
            data = await websocket.receive_text()
            if data == "ping":
                await websocket.send_text("pong")
    except WebSocketDisconnect:
        ws_manager.disconnect(client_id)


@app.websocket("/ws/{client_id}")
async def websocket_endpoint(websocket: WebSocket, client_id: str):
    await ws_manager.connect(websocket, client_id)
    try:
        while True:
            data = await websocket.receive_text()
            if data == "ping":
                await websocket.send_text("pong")
    except WebSocketDisconnect:
        ws_manager.disconnect(client_id)


@app.get("/")
def root():
    return {
        "message": "MSS AI PPT backend is running.",
        "endpoints": [
            "/health",
            "/templates",
            "/inputs",
            "/generate",
            "/upload-excel",  # Excel file upload (NEW)
            "/preview",
            "/download",
            "/cleanup",
            "/docs"
        ],
    }


@app.get("/health")
def health():
    """Health check endpoint with LLM concurrency status."""
    # Get current semaphore state
    # Note: _value is the internal counter in asyncio.Semaphore
    available_slots = llm_semaphore._value if hasattr(llm_semaphore, '_value') else MAX_CONCURRENT_LLM_REQUESTS

    return {
        "status": "ok",
        "llm_concurrency": {
            "max_concurrent_requests": MAX_CONCURRENT_LLM_REQUESTS,
            "available_slots": available_slots,
            "active_requests": MAX_CONCURRENT_LLM_REQUESTS - available_slots,
            "is_saturated": available_slots == 0
        }
    }


@app.get("/templates")
def list_templates():
    return service.template_repo.list_templates()


@app.get("/inputs")
def list_inputs():
    return service.list_inputs()


@app.post("/generate")
async def generate(req: GenerateRequest):
    """Generate a report with OpenAI concurrency limiting.

    This endpoint uses a semaphore to limit concurrent LLM requests,
    preventing API rate limit errors when multiple users generate reports simultaneously.
    """
    # Register WebSocket connection if client_id provided
    if req.client_id and req.session_id:
        ws_manager.register_session(req.session_id, req.client_id)
        # Send initial progress
        await ws_manager.send_progress_update(
            req.session_id, 0, "请求已接收,正在排队...",
            {"template": req.template_id}
        )

    logger.info(f"=== Generate Request: input_id={req.input_id}, template_id={req.template_id}, use_mock={req.use_mock}, session_id={req.session_id} ===")

    # Check if we're using mock mode (no LLM call needed)
    if req.use_mock:
        # Mock mode: no LLM call, no need for semaphore
        try:
            # Send progress: starting mock generation
            if req.client_id and req.session_id:
                await ws_manager.send_progress_update(
                    req.session_id, 20, "快速生成模式：正在生成报告..."
                )

            # Run the synchronous generate() in a thread pool
            loop = asyncio.get_running_loop()
            result = await loop.run_in_executor(
                None,
                lambda: service.generate(
                    req.input_id,
                    req.template_id,
                    use_mock=req.use_mock,
                    session_id=req.session_id
                )
            )

            logger.info(f"✓ Mock generation successful: {result.get('job_id')}")

            # Send progress: generation complete
            if req.client_id and req.session_id:
                await ws_manager.send_progress_update(
                    req.session_id, 90, "报告生成完成，准备预览..."
                )

            # Send completion notification
            if req.client_id and req.session_id:
                await ws_manager.send_completion(
                    req.session_id,
                    result,
                    success=True
                )

            return JSONResponse(result)
        except Exception as e:
            logger.exception(f"✗ Mock generation failed: {e}")

            # Send failure notification
            if req.client_id and req.session_id:
                await ws_manager.send_completion(
                    req.session_id,
                    {"error": str(e)},
                    success=False
                )

            raise HTTPException(status_code=500, detail=str(e))

    # Real LLM mode: acquire semaphore to limit concurrency
    waiting_count = 0
    if hasattr(llm_semaphore, '_waiters') and llm_semaphore._waiters is not None:
        waiting_count = len(llm_semaphore._waiters)
    logger.info(f"🔄 Waiting for LLM slot (current waiting: {waiting_count})")

    async with llm_semaphore:
        logger.info(f"✅ Acquired LLM slot, starting generation...")

        # Send progress: started generation (10%)
        if req.client_id and req.session_id:
            await ws_manager.send_progress_update(
                req.session_id, 10, "已获得处理槽位,开始生成报告..."
            )
        try:
            # Send progress: loading template (20%)
            if req.client_id and req.session_id:
                await ws_manager.send_progress_update(
                    req.session_id, 20, "正在加载模板..."
                )

            # Send progress: AI generation (40%)
            if req.client_id and req.session_id:
                await ws_manager.send_progress_update(
                    req.session_id, 40, "正在调用AI生成内容..."
                )

            # Run the synchronous generate() in a thread pool to avoid blocking
            loop = asyncio.get_running_loop()
            result = await loop.run_in_executor(
                None,  # Use default executor
                lambda: service.generate(
                    req.input_id,
                    req.template_id,
                    use_mock=req.use_mock,
                    session_id=req.session_id
                )
            )
            logger.info(f"✓ Generation successful: {result.get('job_id')}")

            # Send progress: rendering complete (80%)
            if req.client_id and req.session_id:
                await ws_manager.send_progress_update(
                    req.session_id, 80, "报告生成完成,准备预览..."
                )

            # Send completion notification (100%)
            if req.client_id and req.session_id:
                await ws_manager.send_completion(
                    req.session_id,
                    result,
                    success=True
                )
            logger.info(f"  - Session ID: {result.get('session_id')}")
            logger.info(f"  - Report path: {result.get('report_path')}")
            logger.info(f"  - Warnings: {len(result.get('warnings', []))}")
            return JSONResponse(result)
        except InputNotFoundError as e:
            logger.error(f"✗ Input not found: {e}")
            raise HTTPException(status_code=404, detail=str(e))
        except TemplateNotFoundError as e:
            logger.error(f"✗ Template not found: {e}")
            raise HTTPException(status_code=404, detail=str(e))
        except LLMGenerationError as e:
            logger.exception(f"✗ LLM generation failed: {e}")
            error_str = str(e)

            # Check if this is a format error after 5 retries
            if "AI响应格式错误" in error_str and "已重试5次" in error_str:
                logger.warning("🔄 AI格式错误已重试5次，自动切换到mock模式...")

                # Send progress notification
                if req.client_id and req.session_id:
                    await ws_manager.send_progress_update(
                        req.session_id, 50, "AI响应格式异常，切换到数据模式生成..."
                    )

                try:
                    # Retry with mock mode
                    loop = asyncio.get_running_loop()
                    result = await loop.run_in_executor(
                        None,
                        lambda: service.generate(
                            req.input_id,
                            req.template_id,
                            use_mock=True,  # Force mock mode
                            session_id=req.session_id
                        )
                    )

                    logger.info(f"✓ Mock fallback successful: {result.get('job_id')}")

                    # Add warning about mock mode
                    if "warnings" not in result:
                        result["warnings"] = []
                    result["warnings"].insert(0, "由于AI响应格式异常，已自动切换到数据模式生成")
                    result["used_mock_fallback"] = True

                    # Send completion
                    if req.client_id and req.session_id:
                        await ws_manager.send_completion(
                            req.session_id,
                            result,
                            success=True
                        )

                    return JSONResponse(result)

                except Exception as fallback_error:
                    logger.exception(f"✗ Mock fallback also failed: {fallback_error}")
                    # Send failure notification
                    if req.client_id and req.session_id:
                        await ws_manager.send_completion(
                            req.session_id,
                            {"error": f"AI和数据模式均失败: {fallback_error}"},
                            success=False
                        )
                    raise HTTPException(status_code=500, detail=f"AI和数据模式均失败: {fallback_error}")

            # Not a format error, handle as before
            cause = getattr(e, "__cause__", None)
            status_code = 503
            code = "LLM_UNAVAILABLE"
            message = "AI 服务暂时不可用"

            if isinstance(cause, RateLimitError):
                status_code = 429
                code = "LLM_RATE_LIMITED"
                message = "AI 服务请求过于频繁"

            return JSONResponse(
                status_code=status_code,
                content={
                    "error": {
                        "code": code,
                        "message": message,
                        "auto_retry_with_mock": True,  # 标记前端自动重试
                        "user_message": "AI 接口暂时不可用，将为您生成包含数据内容的报告（AI 生成部分使用占位符）。您可以稍后重新生成完整版本。",
                        "details": {
                            "attempts": 4,
                            "model": "GLM4.7",
                            "last_error_type": type(cause).__name__ if cause else type(e).__name__,
                            "last_error_message": str(cause) if cause else str(e),
                        },
                    }
                },
            )
        except Exception as e:
            logger.exception(f"✗ Generation failed with exception: {e}")

            # Send failure notification
            if req.client_id and req.session_id:
                await ws_manager.send_completion(
                    req.session_id,
                    {"error": str(e)},
                    success=False
                )

            raise HTTPException(status_code=500, detail=str(e))
        finally:
            logger.info(f"🔓 Released LLM slot")


@app.post("/rewrite")
def rewrite(req: RewriteRequest):
    try:
        result = service.rewrite(req.job_id, req.slide_key, req.new_content)
        return JSONResponse(result)
    except SlideSpecNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/preview")
def preview(job_id: str, regenerate_if_missing: bool = True):
    try:
        result = service.preview(job_id, regenerate_if_missing=regenerate_if_missing)
        return JSONResponse(result)
    except SlideSpecNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/download")
def download(job_id: str, regenerate_if_missing: bool = True):
    try:
        report_path = service.get_report_path(job_id, regenerate_if_missing=regenerate_if_missing)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        download_name = f"{ts}_{report_path.name}"
        return FileResponse(
            path=report_path,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            filename=download_name,
        )
    except SlideSpecNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/logs")
def logs(limit: int = 100):
    try:
        content = service.read_logs(limit=limit)
        return JSONResponse({"lines": content.splitlines() if content else []})
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/cleanup")
def cleanup_sessions(max_age_hours: int = 24):
    """Clean up old session directories.

    Args:
        max_age_hours: Maximum age in hours before cleanup (default: 24)

    Returns:
        Number of sessions cleaned up
    """
    try:
        cleaned_count = service.cleanup_old_sessions(max_age_hours)
        logger.info(f"Cleaned up {cleaned_count} old sessions")
        return JSONResponse({
            "cleaned_count": cleaned_count,
            "max_age_hours": max_age_hours
        })
    except Exception as e:
        logger.exception(f"✗ Cleanup failed: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ==================== Excel Upload Endpoint ====================

# Configuration for Excel upload
ALLOWED_EXTENSIONS = {'.xlsx'}
FORBIDDEN_EXTENSIONS = {'.xlsm', '.xls', '.xlsb', '.csv'}
ALLOWED_MIME_TYPES = {
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'application/zip',
}
MAX_EXCEL_SIZE = 10 * 1024 * 1024  # 10MB
EXCEL_CHUNK_SIZE = 8192  # 8KB


def extract_excel_data(excel_path: str) -> dict:
    """Extract data from Excel file following agreed format.

    Expected Excel format:
    - Sheet name: "数据表"
    - A2: customer_name
    - B2:C2: period (start, end)
    - D2: total alerts
    - E2:H2: alerts by severity
    - A5:B10: alert categories
    - A15:E20: incident details
    """
    try:
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)

        if '数据表' not in wb.sheetnames:
            raise ValueError(f"工作表 '数据表' 不存在。找到的工作表: {', '.join(wb.sheetnames)}")

        ws = wb['数据表']

        # Extract basic info with validation
        customer_name = ws['A2'].value
        if not customer_name:
            raise ValueError("必填字段 'customer_name' (A2) 为空")

        period_start = ws['B2'].value
        if not period_start:
            raise ValueError("必填字段 'period.start' (B2) 为空")

        period_end = ws['C2'].value
        if not period_end:
            raise ValueError("必填字段 'period.end' (C2) 为空")

        data = {
            "customer_name": str(customer_name).strip(),
            "period": {
                "start": str(period_start),
                "end": str(period_end)
            },
            "alerts": {
                "total": ws['D2'].value or 0,
                "by_severity": {
                    "high": ws['E2'].value or 0,
                    "medium": ws['F2'].value or 0,
                    "low": ws['G2'].value or 0,
                    "info": ws['H2'].value or 0
                }
            },
            "top_alert_categories": [],
            "top_incidents": []
        }

        # Extract alert categories (A5:B10)
        for row in range(5, 11):
            category = ws[f'A{row}'].value
            count = ws[f'B{row}'].value
            if category and count is not None:
                data["top_alert_categories"].append({
                    "category": str(category).strip(),
                    "count": int(count) if isinstance(count, (int, float)) else 0
                })

        # Extract incidents (A15:E20)
        for row in range(15, 21):
            incident_id = ws[f'A{row}'].value
            if incident_id:
                data["top_incidents"].append({
                    "id": str(incident_id).strip(),
                    "type": str(ws[f'B{row}'].value or "").strip(),
                    "severity": str(ws[f'C{row}'].value or "").strip(),
                    "status": str(ws[f'D{row}'].value or "").strip(),
                    "description": str(ws[f'E{row}'].value or "").strip()
                })

        wb.close()
        logger.info(f"✅ Excel parsed: {len(data['top_alert_categories'])} categories, {len(data['top_incidents'])} incidents")
        return data

    except openpyxl.utils.exceptions.InvalidFileException as e:
        raise ValueError(f"无效的Excel文件格式: {str(e)}")
    except Exception as e:
        raise ValueError(f"Excel解析失败: {str(e)}")


@app.post("/upload-excel")
async def upload_excel(file: UploadFile = File(...)):
    """Upload Excel file and parse to JSON.

    Security features:
    - Only accepts .xlsx format (no macros)
    - File size limit: 10MB
    - MIME type validation
    - Session isolation storage

    Returns:
        session_id, filename, preview data
    """
    logger.info(f"📥 Upload request: filename={file.filename}, content_type={file.content_type}")

    # 1. Validate file extension
    file_ext = Path(file.filename).suffix.lower()

    if file_ext not in ALLOWED_EXTENSIONS:
        if file_ext in FORBIDDEN_EXTENSIONS:
            format_tips = {
                '.xlsm': '此格式可能包含宏代码，请使用Excel另存为 .xlsx 格式',
                '.xls': '旧版Excel格式，请使用Excel 2007+另存为 .xlsx 格式',
                '.xlsb': '二进制Excel格式，请另存为 .xlsx 格式',
                '.csv': 'CSV格式不支持，请使用Excel打开并另存为 .xlsx 格式'
            }
            detail_msg = format_tips.get(file_ext, '请转换为 .xlsx 格式')
            raise HTTPException(400, detail={
                "error": "FORBIDDEN_FORMAT",
                "message": f"禁止上传 {file_ext} 格式文件",
                "detail": detail_msg,
                "allowed_formats": list(ALLOWED_EXTENSIONS)
            })
        else:
            raise HTTPException(400, detail={
                "error": "UNSUPPORTED_FORMAT",
                "message": f"不支持的文件格式 {file_ext}，仅允许 .xlsx",
                "allowed_formats": list(ALLOWED_EXTENSIONS)
            })

    # 2. Validate MIME type
    if file.content_type not in ALLOWED_MIME_TYPES:
        logger.warning(f"MIME mismatch: {file.filename} - {file.content_type}")
        raise HTTPException(400, detail={
            "error": "MIME_TYPE_MISMATCH",
            "message": "文件类型验证失败，文件可能被伪装",
            "detail": f"检测到 {file.content_type}，但期望 .xlsx 格式"
        })

    # 3. Generate session and save file
    session_id = service.session_manager.generate_session_id()
    session_dir = config.SESSIONS_DIR / session_id
    session_dir.mkdir(parents=True, exist_ok=True)

    file_path = session_dir / "uploaded.xlsx"
    file_size = 0

    try:
        # Save file with size check
        with file_path.open("wb") as f:
            while chunk := await file.read(EXCEL_CHUNK_SIZE):
                file_size += len(chunk)
                if file_size > MAX_EXCEL_SIZE:
                    f.close()
                    file_path.unlink()
                    raise HTTPException(413, detail={
                        "error": "FILE_TOO_LARGE",
                        "message": f"文件大小超过限制 {MAX_EXCEL_SIZE / 1024 / 1024}MB",
                        "actual_size_mb": round(file_size / 1024 / 1024, 2)
                    })
                f.write(chunk)

        logger.info(f"✅ File saved: {file_path} ({file_size / 1024:.2f} KB)")

        # 4. Parse Excel
        try:
            input_data = extract_excel_data(str(file_path))
        except ValueError as e:
            file_path.unlink()
            raise HTTPException(400, detail={
                "error": "DATA_VALIDATION_FAILED",
                "message": str(e),
                "suggestion": "请检查Excel文件是否按照模板格式填写"
            })

        # 5. Save JSON
        json_path = session_dir / "input.json"
        with json_path.open("w", encoding="utf-8") as f:
            json.dump(input_data, f, ensure_ascii=False, indent=2)

        logger.info(f"✅ JSON saved: {json_path}")

        # 6. Return success response
        return JSONResponse({
            "status": "success",
            "session_id": session_id,
            "filename": file.filename,
            "file_size_mb": round(file_size / 1024 / 1024, 2),
            "files": {
                "excel": str(file_path.name),
                "json": str(json_path.name)
            },
            "preview": {
                "customer_name": input_data.get("customer_name"),
                "period": input_data.get("period"),
                "alerts": {
                    "total": input_data.get("alerts", {}).get("total"),
                    "high": input_data.get("alerts", {}).get("by_severity", {}).get("high")
                },
                "categories_count": len(input_data.get("top_alert_categories", [])),
                "incidents_count": len(input_data.get("top_incidents", []))
            },
            "next_steps": {
                "description": "使用此session_id调用 /generate 端点生成报告",
                "example_request": {
                    "method": "POST",
                    "endpoint": "/generate",
                    "body": {
                        "input_id": "custom",
                        "template_id": "mss_executive_v2",
                        "session_id": session_id
                    }
                }
            }
        })

    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Unexpected error during Excel upload")
        if file_path.exists():
            try:
                file_path.unlink()
            except:
                pass
        raise HTTPException(500, detail={
            "error": "INTERNAL_ERROR",
            "message": f"服务器内部错误: {str(e)}"
        })


# Startup: Clean up old sessions on server start
@app.on_event("startup")
async def startup_cleanup():
    """Clean up old sessions when server starts."""
    try:
        cleaned_count = service.cleanup_old_sessions(max_age_hours=24)
        if cleaned_count > 0:
            logger.info(f"🧹 Startup cleanup: removed {cleaned_count} old sessions")
    except Exception as e:
        logger.warning(f"⚠️ Startup cleanup failed: {e}")


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8000)
