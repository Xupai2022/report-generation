"""Excel file upload and processing module.

This module handles:
- File validation (format, size, MIME type)
- Excel data extraction to JSON
- Session-based storage for concurrent requests

Current parser target:
- `data.xlsx` workbook layout with sheet `数据统计`
"""

from __future__ import annotations

import json
import logging
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, Dict, List, Set

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from mss_ai_ppt_sample_assets.backend.exceptions import (
    DataValidationError,
    FileValidationError,
)

logger = logging.getLogger(__name__)


class ExcelValidator:
    """Validates Excel file uploads for security and format compliance."""

    ALLOWED_EXTENSIONS: Set[str] = {".xlsx"}
    FORBIDDEN_EXTENSIONS: Set[str] = {".xlsm", ".xls", ".xlsb", ".csv"}
    ALLOWED_MIME_TYPES: Set[str] = {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/zip",
    }

    FORMAT_TIPS = {
        ".xlsm": "This format may contain macros. Please convert to .xlsx.",
        ".xls": "Legacy Excel format. Please save as .xlsx.",
        ".xlsb": "Binary Excel format. Please save as .xlsx.",
        ".csv": "CSV is not supported. Please save as .xlsx.",
    }

    def __init__(self, max_size_mb: int = 10):
        self.max_size_bytes = max_size_mb * 1024 * 1024

    def validate_extension(self, filename: str) -> None:
        file_ext = Path(filename).suffix.lower()
        if file_ext not in self.ALLOWED_EXTENSIONS:
            if file_ext in self.FORBIDDEN_EXTENSIONS:
                detail_msg = self.FORMAT_TIPS.get(file_ext, "Please convert to .xlsx.")
                raise FileValidationError(
                    filename=filename,
                    reason=f"Forbidden file extension {file_ext}: {detail_msg}",
                )
            raise FileValidationError(
                filename=filename,
                reason=f"Unsupported file extension {file_ext}, only .xlsx is allowed.",
            )

    def validate_mime_type(self, filename: str, content_type: str) -> None:
        if content_type not in self.ALLOWED_MIME_TYPES:
            logger.warning("MIME mismatch: %s - %s", filename, content_type)
            raise FileValidationError(
                filename=filename,
                reason=f"MIME verification failed: got {content_type}, expected .xlsx MIME.",
            )

    def validate_size(self, filename: str, size_bytes: int) -> None:
        if size_bytes > self.max_size_bytes:
            max_mb = self.max_size_bytes / 1024 / 1024
            actual_mb = round(size_bytes / 1024 / 1024, 2)
            raise FileValidationError(
                filename=filename,
                reason=f"File size {actual_mb}MB exceeds {max_mb}MB limit.",
            )


class ExcelDataExtractor:
    """Extracts structured data from Excel files."""

    DEFAULT_TEMPLATE_ID = "mss_classic_ops"
    CLASSIC_REQUIRED_SHEET = "数据统计"

    @classmethod
    def _get_template_config(cls, template_id: str | None = None) -> Dict[str, Any]:
        resolved_template_id = (template_id or cls.DEFAULT_TEMPLATE_ID).strip() or cls.DEFAULT_TEMPLATE_ID
        config = cls.TEMPLATE_EXTRACTORS.get(resolved_template_id)
        if not config:
            supported_templates = ", ".join(sorted(cls.TEMPLATE_EXTRACTORS.keys()))
            raise DataValidationError(
                field="template_id",
                message=(
                    f"No Excel extractor registered for template '{resolved_template_id}'. "
                    f"Supported templates: {supported_templates}."
                ),
                template_id=resolved_template_id,
            )
        return config

    @staticmethod
    def _unwrap_cell(value: Any) -> tuple[Any, str]:
        if hasattr(value, "value") and hasattr(value, "number_format"):
            return value.value, str(value.number_format or "")
        return value, ""

    @staticmethod
    def _has_value(value: Any) -> bool:
        raw_value, _ = ExcelDataExtractor._unwrap_cell(value)
        if raw_value is None:
            return False
        if isinstance(raw_value, str):
            return bool(raw_value.strip())
        return True

    @staticmethod
    def _decimal_places_from_number_format(number_format: str) -> int | None:
        if not number_format:
            return None

        fmt = number_format.split(";")[0].strip()
        if not fmt or fmt.lower() == "general":
            return None

        # Remove literal/text parts and Excel fill/alignment directives.
        fmt = re.sub(r'"[^"]*"', "", fmt)
        fmt = re.sub(r"\\.", "", fmt)
        fmt = re.sub(r"_.", "", fmt)
        fmt = re.sub(r"\*.", "", fmt)

        decimal_match = re.search(r"\.([0#]+)", fmt)
        if decimal_match:
            return len(decimal_match.group(1))
        if re.search(r"[0#]", fmt):
            return 0
        return None

    @staticmethod
    def _format_numeric_for_text(value: float, number_format: str) -> str:
        decimals = ExcelDataExtractor._decimal_places_from_number_format(number_format)
        is_percent = "%" in number_format
        normalized = ExcelDataExtractor._normalize_number(value)

        if is_percent:
            if decimals is None:
                decimals = 2
            decimals = min(decimals, 2)
            percent_value = ExcelDataExtractor._normalize_number(float(normalized) * 100)
            if isinstance(percent_value, int):
                return f"{percent_value}%"
            return f"{percent_value:.{decimals}f}%"

        if decimals is not None:
            decimals = min(decimals, 2)
            if decimals == 0:
                return str(int(round(float(normalized))))
            return f"{float(normalized):.{decimals}f}"

        if isinstance(normalized, int):
            return str(normalized)
        return f"{float(normalized):.2f}".rstrip("0").rstrip(".")

    @staticmethod
    def _normalize_number(value: Any) -> Any:
        """Normalize numeric value: keep integers, round non-integers to max 2 decimals."""
        if isinstance(value, bool):
            return value
        if not isinstance(value, (int, float)):
            return value

        rounded = round(float(value), 2)
        if rounded.is_integer():
            return int(rounded)
        return rounded

    @staticmethod
    def _format_numbers_for_output(value: Any) -> Any:
        """Recursively format numeric payload for JSON output.

        Rule:
        - Integers remain numeric.
        - Non-integer floats are rendered as fixed 2-decimal strings.
        """
        if isinstance(value, dict):
            return {k: ExcelDataExtractor._format_numbers_for_output(v) for k, v in value.items()}
        if isinstance(value, list):
            return [ExcelDataExtractor._format_numbers_for_output(v) for v in value]
        if isinstance(value, bool):
            return value
        if isinstance(value, int):
            return value
        if isinstance(value, float):
            rounded = round(value, 2)
            if rounded.is_integer():
                return int(rounded)
            return f"{rounded:.2f}"
        return value

    @staticmethod
    def _to_chart_number(value: Any) -> Any:
        """Convert numeric-like values used by chart payloads back to numbers."""
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return ExcelDataExtractor._normalize_number(value)
        if not isinstance(value, str):
            return value

        text = value.strip()
        if text == "":
            return value
        try:
            if text.endswith("%"):
                return ExcelDataExtractor._normalize_number(float(text[:-1]) / 100)
            return ExcelDataExtractor._normalize_number(float(text.replace(",", "")))
        except Exception:
            return value

    @staticmethod
    def _normalize_chart_payload_numbers(value: Any) -> Any:
        """Normalize chart metric arrays to numeric types for downstream chart rendering."""
        numeric_array_keys = {
            "avg_response_minutes",
            "external_attacks",
            "malicious_outbound",
            "alert_counts",
            "valid_incident_counts",
            "risk_host_counts",
            "internal_lateral_attack_counts",
            "attack_counts",
            "defense_rates",
            "values",
        }

        if isinstance(value, list):
            return [ExcelDataExtractor._normalize_chart_payload_numbers(v) for v in value]

        if not isinstance(value, dict):
            return value

        normalized: Dict[str, Any] = {}
        for key, item in value.items():
            if key in numeric_array_keys and isinstance(item, list):
                normalized[key] = [ExcelDataExtractor._to_chart_number(v) for v in item]
                continue

            if key == "series" and isinstance(item, list):
                next_series = []
                for series_item in item:
                    if not isinstance(series_item, dict):
                        next_series.append(series_item)
                        continue
                    next_item = dict(series_item)
                    raw_values = next_item.get("values")
                    if isinstance(raw_values, list):
                        next_item["values"] = [ExcelDataExtractor._to_chart_number(v) for v in raw_values]
                    next_series.append(next_item)
                normalized[key] = next_series
                continue

            normalized[key] = ExcelDataExtractor._normalize_chart_payload_numbers(item)

        return normalized

    @staticmethod
    def _to_text(value: Any, default: str = "") -> str:
        raw_value, number_format = ExcelDataExtractor._unwrap_cell(value)

        if raw_value is None:
            return default
        if isinstance(raw_value, (datetime, date)):
            return raw_value.strftime("%Y-%m-%d")
        if isinstance(raw_value, (int, float)) and not isinstance(raw_value, bool):
            return ExcelDataExtractor._format_numeric_for_text(float(raw_value), number_format)

        text = str(raw_value).strip()
        return text if text else default

    @staticmethod
    def _format_wan_text(value: float) -> str:
        if abs(value) >= 10000:
            return f"{value / 10000:.2f}万"
        if float(value).is_integer():
            return str(int(value))
        return f"{value:.2f}".rstrip("0").rstrip(".")

    @staticmethod
    def _to_text_wan(value: Any, default: str = "") -> str:
        raw_value, _ = ExcelDataExtractor._unwrap_cell(value)

        if raw_value is None:
            return default
        if isinstance(raw_value, (int, float)) and not isinstance(raw_value, bool):
            return ExcelDataExtractor._format_wan_text(float(raw_value))

        text = str(raw_value).strip()
        if text == "":
            return default
        try:
            numeric = float(text.replace(",", ""))
        except Exception:
            return text
        return ExcelDataExtractor._format_wan_text(numeric)

    @staticmethod
    def _to_number(value: Any, default: float = 0) -> float:
        raw_value, _ = ExcelDataExtractor._unwrap_cell(value)

        if raw_value is None:
            return default
        if isinstance(raw_value, (int, float)):
            return ExcelDataExtractor._normalize_number(raw_value)

        text = str(raw_value).strip()
        if text in {"", "None", "#DIV/0!", "#N/A"}:
            return default

        if text.endswith("%"):
            try:
                return ExcelDataExtractor._normalize_number(float(text[:-1]) / 100)
            except Exception:
                return default

        try:
            if "." in text:
                return ExcelDataExtractor._normalize_number(float(text))
            return ExcelDataExtractor._normalize_number(float(int(text)))
        except Exception:
            return default

    @staticmethod
    def _to_pct(value: Any) -> str:
        raw_value, number_format = ExcelDataExtractor._unwrap_cell(value)
        if isinstance(raw_value, (int, float)) and "%" in number_format:
            decimals = ExcelDataExtractor._decimal_places_from_number_format(number_format)
            if decimals is None:
                decimals = 2
            return f"{raw_value * 100:.{decimals}f}%"

        n = ExcelDataExtractor._to_number(value, 0)
        if n <= 1:
            return f"{round(n * 100, 2)}%"
        return f"{round(n, 2)}%"

    @staticmethod
    def _put_text(target: Dict[str, Any], key: str, raw_value: Any) -> None:
        if ExcelDataExtractor._has_value(raw_value):
            target[key] = ExcelDataExtractor._to_text(raw_value)

    @staticmethod
    def _put_text_wan(target: Dict[str, Any], key: str, raw_value: Any) -> None:
        if ExcelDataExtractor._has_value(raw_value):
            target[key] = ExcelDataExtractor._to_text_wan(raw_value)

    @staticmethod
    def _put_pct(target: Dict[str, Any], key: str, raw_value: Any) -> None:
        if ExcelDataExtractor._has_value(raw_value):
            target[key] = ExcelDataExtractor._to_pct(raw_value)

    @staticmethod
    def _add_section(root: Dict[str, Any], section_key: str, section: Dict[str, Any]) -> None:
        if section:
            root[section_key] = section

    @staticmethod
    def _read_month_columns(ws, header_row: int, col_start: int, col_end: int) -> List[int]:
        cols: List[int] = []
        for c in range(col_start, col_end + 1):
            if ExcelDataExtractor._has_value(ws.cell(header_row, c).value):
                cols.append(c)
        return cols

    @staticmethod
    def _read_labeled_pairs(ws, start_row: int, end_row: int, label_col: int, value_col: int) -> Dict[str, List[Any]]:
        labels: List[str] = []
        values: List[float] = []
        for r in range(start_row, end_row + 1):
            raw_label = ws.cell(r, label_col).value
            raw_value = ws.cell(r, value_col).value
            if not ExcelDataExtractor._has_value(raw_label) and not ExcelDataExtractor._has_value(raw_value):
                continue
            labels.append(ExcelDataExtractor._to_text(raw_label))
            values.append(ExcelDataExtractor._to_number(raw_value, 0))
        return {"labels": labels, "values": values}

    @staticmethod
    def _remove_ai_generated_fields(data: Dict[str, Any]) -> Dict[str, Any]:
        """Remove keys mapped to ai_generate placeholders from classic descriptor."""
        descriptor_dir = Path(__file__).resolve().parent.parent / "data" / "templates"
        descriptor = None

        for path in descriptor_dir.rglob("*_descriptor.json"):
            try:
                content = json.loads(path.read_text(encoding="utf-8-sig"))
            except Exception:
                continue
            if content.get("template_id") == "mss_classic_ops":
                descriptor = content
                break

        if not descriptor:
            return data

        for slide in descriptor.get("slides", []):
            slide_key = slide.get("slide_key")
            if not slide_key:
                continue

            section = data.get(slide_key)
            if not isinstance(section, dict):
                continue

            ai_tokens = {
                ph.get("token")
                for ph in slide.get("placeholders", [])
                if ph.get("ai_generate") and ph.get("token")
            }
            for token in ai_tokens:
                section.pop(token, None)

            if not section:
                data.pop(slide_key, None)

        return data

    @staticmethod
    def _extract_classic_ops(ws) -> Dict[str, Any]:
        output: Dict[str, Any] = {
            "schema_version": "1.0",
            "template_id": "mss_classic_ops",
        }

        raw_period_start = ws["L1"]
        raw_period_end = ws["M1"]
        period: Dict[str, Any] = {}

        if ExcelDataExtractor._has_value(raw_period_start):
            start = ExcelDataExtractor._to_text(raw_period_start)
            period["start"] = start
            period["start_month"] = start[:7] if len(start) >= 7 else start

        if ExcelDataExtractor._has_value(raw_period_end):
            end = ExcelDataExtractor._to_text(raw_period_end)
            period["end"] = end
            period["end_month"] = end[:7] if len(end) >= 7 else end

        if period:
            output["period"] = period

        cover: Dict[str, Any] = {}
        if period.get("start"):
            cover["period_start"] = period["start"]
        if period.get("end"):
            cover["period_end"] = period["end"]
        ExcelDataExtractor._add_section(output, "cover", cover)

        architecture: Dict[str, Any] = {}
        ExcelDataExtractor._put_text(architecture, "AF_count", ws["D3"])
        ExcelDataExtractor._put_text(architecture, "STA_count", ws["D4"])
        ExcelDataExtractor._put_text(architecture, "EDR_count", ws["D5"])
        ExcelDataExtractor._put_text(architecture, "TSS_count", ws["D6"])
        ExcelDataExtractor._add_section(output, "Architecture", architecture)

        deliverables: Dict[str, Any] = {}
        deliverables_map = {
            "first_quarter_date_range": "D8",
            "second_quarter_date_range": "F8",
            "third_quarter_date_range": "H8",
            "fourth_quarter_date_range": "J8",
            "project_launch_ppt": "D9",
            "initial_analysis_and_disposal_report": "D10",
            "vulnerability_management_evidence_report": "D11",
            "vulnerability_list": "D12",
            "service_asset_table": "D13",
            "incident_tracking_table": "D14",
            "emergency_response_report": "D15",
            "exposure_surface_analysis_report": "D16",
            "security_operations_report_weekly": "G9",
            "comprehensive_analysis_report_monthly": "G10",
            "security_operations_report_quarterly": "G11",
            "important_holiday_network_security_work_report": "G13",
            "security_trends_semi_monthly_report": "G17",
            "phishing_scenario_security_poster": "G18",
        }
        for token, addr in deliverables_map.items():
            ExcelDataExtractor._put_text(deliverables, token, ws[addr])
        ExcelDataExtractor._add_section(output, "deliverables", deliverables)

        coverage_summary: Dict[str, Any] = {}
        ExcelDataExtractor._put_pct(coverage_summary, "AF_protection_path_coverage_rate", ws["L22"])
        ExcelDataExtractor._put_pct(coverage_summary, "probe_traffic_monitoring_coverage_rate", ws["L23"])
        ExcelDataExtractor._put_pct(coverage_summary, "AES_installation_coverage_rate", ws["L24"])
        coverage_map = {
            "cybersecurity_incident": "C21",
            "proactive_protection": "D21",
            "incident_count1": "E21",
            "closure_rate1": "F21",
            "incident_count2": "G21",
            "closure_rate2": "H21",
            "response_time": "I21",
            "alert_automatic_analysis_ratio": "L26",
            "alert_manual_analysis_ratio": "L25",
            "mss_risk_asset": "J22",
            "nonmss_risk_asset": "J27",
            "alert_mss": "J23",
            "incident_mss": "J24",
            "handling_rate_mss": "J25",
            "response_time_mss": "J26",
            "alert_nonmss": "J28",
            "incident_nonmss": "J29",
            "handling_rate_nonmss": "J30",
            "response_time_nonmss": "J31",
            "business_system": "D28",
            "server_asset": "D29",
            "pc_asset": "D30",
            "number_of_security_logs": "G22",
            "log_noise_reduction_rate": "G23",
            "number_of_security_alerts": "G24",
            "alert_reduction_rate": "G25",
            "number_of_security_incidents": "G26",
            "number_of_risk_assets": "G27",
            "AF_count": "D22",
            "STA_count": "D23",
            "EDR_count": "D24",
            "TSS_count": "D25",
        }
        wan_tokens = {
            "number_of_security_logs",
            "number_of_security_alerts",
            "number_of_security_incidents",
            "alert_mss",
            "incident_mss",
            "alert_nonmss",
            "incident_nonmss",
        }
        for token, addr in coverage_map.items():
            if token in wan_tokens:
                ExcelDataExtractor._put_text_wan(coverage_summary, token, ws[addr])
            else:
                ExcelDataExtractor._put_text(coverage_summary, token, ws[addr])
        ExcelDataExtractor._add_section(output, "coverage_summary", coverage_summary)

        protection_overview: Dict[str, Any] = {}
        protection_map = {
            "AF_external_attack_blocks": "D35",
            "EDR_endpoint_risk_count": "D36",
            "policy_check_and_optimization_count": "D37",
            "high_risk_exploitable_vulnerability_protection_rate": "C34",
            "high_risk_exploitable_vulnerability_protection_count": "D38",
            "exposure_surface_scan_count": "D39",
            "vulnerability_and_scan_count": "D40",
            "service_asset_count": "D41",
            "PC_asset_count": "D42",
            "high_risk_exploitable_vulnerability_count": "D43",
            "exposure_surface_risk_count": "D44",
            "closed_loop_incident_ticket_count": "D45",
            "closed_loop_vulnerability_ticket_count": "D46",
            "operational_weekly_report_count": "D47",
            "operational_monthly_report_count": "D48",
            "alert_analysis_rate": "D34",
            "alert_average_response_time": "E34",
            "real_time_threat_tracking_rate": "F34",
            "XDR_security_log_count": "G35",
            "XDR_security_alert_total_count": "G36",
            "XDR_security_incident_count": "G37",
            "MSS_pushed_event_count": "G38",
            "MSS_event_average_push_duration": "G39",
            "XDR_automated_processing_event_count": "G40",
            "MSS_processed_event_count": "G41",
            "emergency_response_event_count": "G42",
            "event_average_response_time": "G43",
        }
        protection_wan_tokens = {
            "AF_external_attack_blocks",
            "XDR_security_log_count",
            "XDR_security_alert_total_count",
            "XDR_security_incident_count",
        }
        for token, addr in protection_map.items():
            if token in protection_wan_tokens:
                ExcelDataExtractor._put_text_wan(protection_overview, token, ws[addr])
            else:
                ExcelDataExtractor._put_text(protection_overview, token, ws[addr])
        ExcelDataExtractor._add_section(output, "protection_overview", protection_overview)

        incident_effectiveness: Dict[str, Any] = {}
        incident_map = {
            "incident_total": "C51",
            "average_response_time": "D51",
            "average_resolution_duration": "E51",
            "event_closed_loop_rate": "F51",
        }
        for token, addr in incident_map.items():
            ExcelDataExtractor._put_text(incident_effectiveness, token, ws[addr])

        response_timeliness = ExcelDataExtractor._read_labeled_pairs(ws, 53, 57, 6, 7)
        if response_timeliness["labels"]:
            incident_effectiveness["response_timeliness"] = response_timeliness

        trend_cols = ExcelDataExtractor._read_month_columns(ws, 60, 3, 12)
        if trend_cols:
            response_trend_months = [ExcelDataExtractor._to_text(ws.cell(60, c)) for c in trend_cols]
            response_trend_values = [
                ExcelDataExtractor._to_number(ws.cell(61, c).value, 0) for c in trend_cols
            ]
            response_trend = {
                "months": response_trend_months,
                "avg_response_minutes": response_trend_values,
            }
            incident_effectiveness["response_trend"] = response_trend

        incident_distribution = ExcelDataExtractor._read_labeled_pairs(ws, 53, 57, 3, 4)
        if incident_distribution["labels"]:
            incident_dist_obj = {
                "categories": incident_distribution["labels"],
                "values": incident_distribution["values"],
            }
            incident_effectiveness["incident_distribution"] = incident_dist_obj

        ExcelDataExtractor._add_section(output, "incident_effectiveness", incident_effectiveness)

        asset_management: Dict[str, Any] = {}
        asset_map = {
            "internal_network_business_area": "B70",
            "external_network_business_area": "E70",
            "internal_network_server_count": "D70",
            "internal_network_network_device_count": "D72",
            "internal_network_IoT_device_count": "D73",
            "internal_network_MSS_service_asset_count": "D74",
            "external_root_domain_asset_count": "G70",
            "external_subdomain_asset_count": "G71",
            "web_asset": "G72",
            "nonweb_asset": "G73",
            "login_endpoint_count": "G74",
            "total_server_assets_count": "D64",
            "PC_assets_count": "D65",
            "internet_IP_domain_count": "D66",
            "internet_exposed_ports_count": "D67",
            "asset_identification_runs": "D68",
        }
        for token, addr in asset_map.items():
            ExcelDataExtractor._put_text(asset_management, token, ws[addr])

        asset_dist = ExcelDataExtractor._read_labeled_pairs(ws, 64, 68, 6, 7)
        if asset_dist["labels"]:
            asset_dist_obj = {
                "categories": asset_dist["labels"],
                "values": asset_dist["values"],
            }
            asset_management["asset_distribution"] = asset_dist_obj

        ExcelDataExtractor._add_section(output, "asset_management", asset_management)

        vulnerability_effectiveness: Dict[str, Any] = {}
        vuln_map = {
            "high_risk_exploitable_vulnerability_count": "C77",
            "closed_loop_external_asset_vulnerability_count": "D77",
            "admin_weak_password_count": "E77",
            "high_risk_exploitable_vulnerability_closure_rate": "F77",
            "scanning": "D40",
        }
        for token, addr in vuln_map.items():
            ExcelDataExtractor._put_text(vulnerability_effectiveness, token, ws[addr])

        vuln_dist = ExcelDataExtractor._read_labeled_pairs(ws, 80, 82, 3, 4)
        if vuln_dist["labels"]:
            vuln_dist_obj = {"categories": vuln_dist["labels"], "values": vuln_dist["values"]}
            vulnerability_effectiveness["vulnerability_distribution"] = vuln_dist_obj

        # Preserve additional P14 columns regardless of current downstream usage.
        vuln_closed_loop_counts = ExcelDataExtractor._read_labeled_pairs(ws, 80, 82, 3, 5)
        if vuln_closed_loop_counts["labels"]:
            vulnerability_effectiveness["closed_loop_counts"] = {
                "categories": vuln_closed_loop_counts["labels"],
                "values": vuln_closed_loop_counts["values"],
            }

        closed_loop_rate_labels: List[str] = []
        closed_loop_rate_values: List[str] = []
        for r in range(80, 83):
            raw_label = ws.cell(r, 3)
            raw_rate = ws.cell(r, 6)
            if not ExcelDataExtractor._has_value(raw_label) and not ExcelDataExtractor._has_value(raw_rate):
                continue
            closed_loop_rate_labels.append(ExcelDataExtractor._to_text(raw_label))
            closed_loop_rate_values.append(ExcelDataExtractor._to_pct(raw_rate))
        if closed_loop_rate_labels:
            vulnerability_effectiveness["closed_loop_rates"] = {
                "categories": closed_loop_rate_labels,
                "values": closed_loop_rate_values,
            }
        ExcelDataExtractor._add_section(output, "vulnerability_effectiveness", vulnerability_effectiveness)

        threat_effectiveness: Dict[str, Any] = {}
        # Always preserve raw monthly rows C97:N100 (12 values each) in input JSON.
        # Keep these under threat_trend for a single coherent trend payload.
        raw_cols = range(3, 15)  # C..N
        threat_trend: Dict[str, Any] = {
            "internal_lateral_attack_counts": [
                ExcelDataExtractor._to_number(ws.cell(97, c).value, 0) for c in raw_cols
            ],
            "alert_counts": [
                ExcelDataExtractor._to_number(ws.cell(98, c).value, 0) for c in raw_cols
            ],
            "valid_incident_counts": [
                ExcelDataExtractor._to_number(ws.cell(99, c).value, 0) for c in raw_cols
            ],
            "risk_host_counts": [
                ExcelDataExtractor._to_number(ws.cell(100, c).value, 0) for c in raw_cols
            ],
        }

        month_cols = ExcelDataExtractor._read_month_columns(ws, 94, 3, 13)
        if month_cols:
            months = [ExcelDataExtractor._to_text(ws.cell(94, c)) for c in month_cols]
            external_attacks = [ExcelDataExtractor._to_number(ws.cell(95, c).value, 0) for c in month_cols]
            malicious_outbound = [ExcelDataExtractor._to_number(ws.cell(96, c).value, 0) for c in month_cols]
            alerts_monthly = [ExcelDataExtractor._to_number(ws.cell(98, c).value, 0) for c in month_cols]
            incidents_monthly = [ExcelDataExtractor._to_number(ws.cell(99, c).value, 0) for c in month_cols]

            threat_trend["months"] = months
            threat_trend["external_attacks"] = external_attacks
            threat_trend["malicious_outbound"] = malicious_outbound

        threat_effectiveness["threat_trend"] = threat_trend

        threat_map = {
            "external_attack_log_count_XDR": "D84",
            "real_time_threat_alert_count": "D85",
            "MSS_threat_ticket_count": "D86",
            "threat_ticket_average_response_time": "D87",
            "security_device_policy_check_count": "D88",
            "optimized_policy_risk_count": "D89",
            "latest_threat_intelligence_count": "D90",
            "latest_threat_impacted_asset_count": "D91",
        }
        for token, addr in threat_map.items():
            ExcelDataExtractor._put_text(threat_effectiveness, token, ws[addr])

        # Page 15 charts:
        # - F/G: 攻击来源地域分布TOP5
        # - I/J: 攻击类型TOP5
        # - L/M: 遭受外部攻击的主机TOP5
        attack_source_region_top5 = ExcelDataExtractor._read_labeled_pairs(ws, 85, 89, 6, 7)
        if attack_source_region_top5["labels"]:
            threat_effectiveness["attack_source_region_top5"] = {
                "categories": attack_source_region_top5["labels"],
                "values": attack_source_region_top5["values"],
            }

        attack_type_top5 = ExcelDataExtractor._read_labeled_pairs(ws, 85, 89, 9, 10)
        if attack_type_top5["labels"]:
            threat_effectiveness["attack_type_top5"] = {
                "categories": attack_type_top5["labels"],
                "values": attack_type_top5["values"],
            }

        externally_attacked_hosts_top5 = ExcelDataExtractor._read_labeled_pairs(ws, 85, 89, 12, 13)
        if externally_attacked_hosts_top5["labels"]:
            threat_effectiveness["externally_attacked_hosts_top5"] = {
                "categories": externally_attacked_hosts_top5["labels"],
                "values": externally_attacked_hosts_top5["values"],
            }
        ExcelDataExtractor._add_section(output, "threat_effectiveness", threat_effectiveness)

        critical_assurance: Dict[str, Any] = {}
        critical_map = {
            "duty_critical": "D102",
            "incident_critical": "D103",
            "availability_assure": "D104",
        }
        for token, addr in critical_map.items():
            ExcelDataExtractor._put_text(critical_assurance, token, ws[addr])

        posture_rows = []
        for r in range(103, 110):
            raw_cat = ws.cell(r, 6).value
            raw_attack = ws.cell(r, 7).value
            raw_defense = ws.cell(r, 8).value
            if (
                not ExcelDataExtractor._has_value(raw_cat)
                and not ExcelDataExtractor._has_value(raw_attack)
                and not ExcelDataExtractor._has_value(raw_defense)
            ):
                continue
            posture_rows.append((raw_cat, raw_attack, raw_defense))

        if posture_rows:
            posture_comparison = {
                "categories": [ExcelDataExtractor._to_text(r[0]) for r in posture_rows],
                "attack_counts": [ExcelDataExtractor._to_number(r[1], 0) for r in posture_rows],
                "defense_rates": [ExcelDataExtractor._to_number(r[2], 0) for r in posture_rows],
            }
            critical_assurance["posture_comparison"] = posture_comparison

        ExcelDataExtractor._add_section(output, "critical_assurance", critical_assurance)

        platform_effectiveness: Dict[str, Any] = {}
        platform_map = {
            "firewall_detected_attack_count": "D111",
            "firewall_automatic_block_rate": "D112",
            "firewall_protection_path_coverage_rate": "D113",
            "AES_detected_endpoint_security_risk_count": "D114",
            "host_anomaly_risk_handling_count": "D115",
            "server_endpoint_coverage_rate": "D116",
            "XDR_total_security_log_count": "D117",
            "aggregated_security_alert_count": "D118",
            "intelligent_security_incident_identification_count": "D119",
            "component_network_connectivity_anomaly_count": "D121",
            "component_log_synchronization_anomaly_count": "D122",
            "component_policy_effectiveness_alert_count": "D123",
            "component_anomaly_automatic_handling_count": "D124",
        }
        platform_wan_tokens = {
            "firewall_detected_attack_count",
            "XDR_total_security_log_count",
            "aggregated_security_alert_count",
            "intelligent_security_incident_identification_count",
        }
        for token, addr in platform_map.items():
            if token in platform_wan_tokens:
                ExcelDataExtractor._put_text_wan(platform_effectiveness, token, ws[addr])
            else:
                ExcelDataExtractor._put_text(platform_effectiveness, token, ws[addr])

        platform_extra_map = {
            "AES_trusted_risk_event_count": "F114",
            "agent_installation_count": "F115",
            "total_asset_count": "F116",
            "XDR_monthly_average_log_count": "F117",
            "XDR_monthly_average_alert_count": "F118",
            "XDR_monthly_average_incident_count": "F119",
        }
        platform_extra_wan_tokens = {
            "XDR_monthly_average_log_count",
            "XDR_monthly_average_alert_count",
            "XDR_monthly_average_incident_count",
        }
        for token, addr in platform_extra_map.items():
            if token in platform_extra_wan_tokens:
                ExcelDataExtractor._put_text_wan(platform_effectiveness, token, ws[addr])
            else:
                ExcelDataExtractor._put_text(platform_effectiveness, token, ws[addr])

        ExcelDataExtractor._add_section(output, "platform_effectiveness", platform_effectiveness)

        output = ExcelDataExtractor._format_numbers_for_output(output)
        output = ExcelDataExtractor._normalize_chart_payload_numbers(output)
        return ExcelDataExtractor._remove_ai_generated_fields(output)

    @staticmethod
    def _extract_classic_ops_2(ws) -> Dict[str, Any]:
        output: Dict[str, Any] = {
            "schema_version": "1.0",
            "template_id": "mss_classic_ops_2",
        }

        raw_period_start = ws["L1"]
        raw_period_end = ws["M1"]
        period: Dict[str, Any] = {}

        if ExcelDataExtractor._has_value(raw_period_start):
            start = ExcelDataExtractor._to_text(raw_period_start)
            period["start"] = start
            period["start_month"] = start[:7] if len(start) >= 7 else start

        if ExcelDataExtractor._has_value(raw_period_end):
            end = ExcelDataExtractor._to_text(raw_period_end)
            period["end"] = end
            period["end_month"] = end[:7] if len(end) >= 7 else end

        if period:
            output["period"] = period

        cover: Dict[str, Any] = {}
        if period.get("start"):
            cover["period_start"] = period["start"]
        if period.get("end"):
            cover["period_end"] = period["end"]
        ExcelDataExtractor._put_text(cover, "company", ws["J1"])
        ExcelDataExtractor._add_section(output, "cover", cover)

        success_metric: Dict[str, Any] = {}
        ExcelDataExtractor._put_text(success_metric, "core_system_1", ws["D3"])
        ExcelDataExtractor._put_text(success_metric, "core_system_2", ws["D4"])
        ExcelDataExtractor._put_text(success_metric, "core_system_3", ws["D5"])
        ExcelDataExtractor._add_section(output, "success_metric", success_metric)

        ensure_result: Dict[str, Any] = {}
        ensure_result_map = {
            "vulnerability_count": "D6",
            "system_1_vulnerability_count": "D7",
            "system_2_vulnerability_count": "D8",
            "system_3_vulnerability_count": "D9",
            "weak_password_count": "D10",
            "latest_vulnerability_intelligence_count": "D11",
            "latest_vulnerability_affected_assets_count": "D12",
            "closed_loop_risk_count": "D14",
            "latest_affected_vulnerability_count": "D15",
            "protected_vulnerability_count": "D16",
            "fixed_vulnerability_count": "D17",
            "vulnerability_exploitation_attempt_count": "D18",
            "policy_optimization_count": "G4",
            "external_threat_alert_count": "G5",
            "average_threat_containment_time": "G6",
            "security_incident_count": "G7",
            "average_incident_response_time": "G8",
            "incident_closure_rate": "G9",
        }
        for token, addr in ensure_result_map.items():
            ExcelDataExtractor._put_text(ensure_result, token, ws[addr])
        ExcelDataExtractor._add_section(output, "ensure_result", ensure_result)

        reduce_security_alert_risk: Dict[str, Any] = {}
        reduce_security_alert_risk_map = {
            "previous_service_period_reported_count": "D20",
            "malicious_external_connection_count": "D21",
            "vulnerability_scan_count": "D22",
            "fixed_vulnerability_count": "D23",
            "protected_vulnerability_count": "D24",
            "reported_count": "D25",
        }
        for token, addr in reduce_security_alert_risk_map.items():
            ExcelDataExtractor._put_text(reduce_security_alert_risk, token, ws[addr])
        ExcelDataExtractor._add_section(output, "reduce_security_alert_risk", reduce_security_alert_risk)

        critical_period_protection: Dict[str, Any] = {}
        critical_period_protection_map = {
            "critical_period_security_log_count": "D27",
            "critical_period_alert_count": "D28",
            "critical_period_incident_count": "D29",
            "containment_rate": "D30",
        }
        for token, addr in critical_period_protection_map.items():
            ExcelDataExtractor._put_text(critical_period_protection, token, ws[addr])
        critical_protection_periods = [
            ExcelDataExtractor._to_text(ws.cell(row, 6))
            for row in range(91, 98)
            if ExcelDataExtractor._has_value(ws.cell(row, 6))
        ]
        if critical_protection_periods:
            critical_period_protection["critical_protection_periods_text"] = "、".join(critical_protection_periods)
        ExcelDataExtractor._add_section(output, "critical_period_protection", critical_period_protection)

        security_level_quantification: Dict[str, Any] = {}
        security_level_quantification_map = {
            "asset_coverage_rate": "D32",
            "threat_containment_rate": "D33",
            "high_risk_vulnerability_protection_rate": "D34",
            "average_analysis_time": "D35",
            "average_response_time": "D36",
            "closed_loop_vulnerability_count": "D37",
            "threat_and_incident_count": "D38",
            "two_senior_security_engineer_cost": "D39",
        }
        for token, addr in security_level_quantification_map.items():
            ExcelDataExtractor._put_text(security_level_quantification, token, ws[addr])
        ExcelDataExtractor._add_section(output, "security_level_quantification", security_level_quantification)

        incident_effectiveness: Dict[str, Any] = {}
        incident_effectiveness_map = {
            "incident_total": "C49",
            "average_response_time": "D49",
            "average_resolution_duration": "E49",
            "event_closed_loop_rate": "F49",
        }
        for token, addr in incident_effectiveness_map.items():
            ExcelDataExtractor._put_text(incident_effectiveness, token, ws[addr])

        response_timeliness_labels: List[str] = []
        response_timeliness_values: List[Any] = []
        for row in (51, 52, 54, 55):
            raw_label = ws.cell(row, 6)
            raw_value = ws.cell(row, 7)
            if not ExcelDataExtractor._has_value(raw_label) and not ExcelDataExtractor._has_value(raw_value):
                continue
            response_timeliness_labels.append(ExcelDataExtractor._to_text(raw_label))
            response_timeliness_values.append(ExcelDataExtractor._to_number(raw_value, 0))
        if response_timeliness_labels:
            incident_effectiveness["response_timeliness"] = {
                "labels": response_timeliness_labels,
                "values": response_timeliness_values,
            }

        trend_cols = range(3, 15)
        response_trend_months = [
            ExcelDataExtractor._to_text(ws.cell(58, col))
            for col in trend_cols
            if ExcelDataExtractor._has_value(ws.cell(58, col))
        ]
        response_trend_values = [
            ExcelDataExtractor._to_number(ws.cell(59, col).value, 0)
            for col in trend_cols
            if ExcelDataExtractor._has_value(ws.cell(58, col))
        ]
        if response_trend_months:
            incident_effectiveness["response_trend"] = {
                "months": response_trend_months,
                "avg_response_minutes": response_trend_values,
            }

        incident_distribution = ExcelDataExtractor._read_labeled_pairs(ws, 51, 55, 3, 4)
        if incident_distribution["labels"]:
            incident_effectiveness["incident_distribution"] = {
                "categories": incident_distribution["labels"],
                "values": incident_distribution["values"],
            }

        ExcelDataExtractor._add_section(output, "incident_effectiveness", incident_effectiveness)

        threat_effectiveness: Dict[str, Any] = {}
        threat_effectiveness_map = {
            "external_attack_log_count_XDR": "D61",
            "real_time_threat_alert_count": "D62",
            "MSS_threat_ticket_count": "D63",
            "threat_ticket_average_response_time": "D64",
            "security_device_policy_check_count": "D65",
            "optimized_policy_risk_count": "D66",
            "latest_threat_intelligence_count": "D67",
            "latest_threat_impacted_asset_count": "D68",
        }
        for token, addr in threat_effectiveness_map.items():
            ExcelDataExtractor._put_text(threat_effectiveness, token, ws[addr])

        attack_source_region_top5 = ExcelDataExtractor._read_labeled_pairs(ws, 62, 66, 6, 7)
        if attack_source_region_top5["labels"]:
            threat_effectiveness["attack_source_region_top5"] = {
                "categories": attack_source_region_top5["labels"],
                "values": attack_source_region_top5["values"],
            }

        attack_type_top5 = ExcelDataExtractor._read_labeled_pairs(ws, 62, 66, 9, 10)
        if attack_type_top5["labels"]:
            threat_effectiveness["attack_type_top5"] = {
                "categories": attack_type_top5["labels"],
                "values": attack_type_top5["values"],
            }

        externally_attacked_hosts_top5 = ExcelDataExtractor._read_labeled_pairs(ws, 62, 66, 12, 13)
        if externally_attacked_hosts_top5["labels"]:
            threat_effectiveness["externally_attacked_hosts_top5"] = {
                "categories": externally_attacked_hosts_top5["labels"],
                "values": externally_attacked_hosts_top5["values"],
            }

        threat_trend_cols = range(3, 15)
        threat_trend_months = [
            ExcelDataExtractor._to_text(ws.cell(71, col))
            for col in threat_trend_cols
            if ExcelDataExtractor._has_value(ws.cell(71, col))
        ]
        if threat_trend_months:
            threat_effectiveness["threat_trend"] = {
                "months": threat_trend_months,
                "external_attacks": [
                    ExcelDataExtractor._to_number(ws.cell(72, col).value, 0)
                    for col in threat_trend_cols
                    if ExcelDataExtractor._has_value(ws.cell(71, col))
                ],
                "malicious_outbound": [
                    ExcelDataExtractor._to_number(ws.cell(73, col).value, 0)
                    for col in threat_trend_cols
                    if ExcelDataExtractor._has_value(ws.cell(71, col))
                ],
                "series_labels": [
                    ExcelDataExtractor._to_text(ws["B72"]),
                    ExcelDataExtractor._to_text(ws["B73"]),
                ],
            }
        ExcelDataExtractor._add_section(output, "threat_effectiveness", threat_effectiveness)

        risk_prevention_work_details: Dict[str, Any] = {}
        risk_prevention_map = {
            "high_risk_exploitable_vulnerability_count": "C80",
            "closed_loop_external_asset_vulnerability_count": "D80",
            "admin_weak_password_count": "E80",
            "high_risk_exploitable_vulnerability_closure_rate": "F80",
            "internet_ip_count": "I79",
            "internet_port_count": "I80",
            "exposed_risk_port_count": "I81",
            "external_network_asset_vulnerability_count": "I82",
            "server_asset_count": "K79",
            "mss_service_asset_count": "K80",
            "total_vulnerability_count": "K81",
            "protected_vulnerability_count": "K82",
            "fixed_vulnerability_count": "K83",
            "vulnerability_scan_count": "N79",
        }
        for token, addr in risk_prevention_map.items():
            ExcelDataExtractor._put_text(risk_prevention_work_details, token, ws[addr])

        vulnerability_distribution = ExcelDataExtractor._read_labeled_pairs(ws, 83, 85, 3, 4)
        if vulnerability_distribution["labels"]:
            risk_prevention_work_details["vulnerability_distribution"] = {
                "categories": vulnerability_distribution["labels"],
                "values": vulnerability_distribution["values"],
            }

        vulnerability_trend_cols = range(3, 15)
        vulnerability_trend_months = [
            ExcelDataExtractor._to_text(ws.cell(86, col))
            for col in vulnerability_trend_cols
            if ExcelDataExtractor._has_value(ws.cell(86, col))
        ]
        if vulnerability_trend_months:
            risk_prevention_work_details["vulnerability_trend"] = {
                "months": vulnerability_trend_months,
                "series": [
                    {
                        "name": ExcelDataExtractor._to_text(ws["B87"]),
                        "values": [
                            ExcelDataExtractor._to_number(ws.cell(87, col).value, 0)
                            for col in vulnerability_trend_cols
                            if ExcelDataExtractor._has_value(ws.cell(86, col))
                        ],
                    },
                    {
                        "name": ExcelDataExtractor._to_text(ws["B88"]),
                        "values": [
                            ExcelDataExtractor._to_number(ws.cell(88, col).value, 0)
                            for col in vulnerability_trend_cols
                            if ExcelDataExtractor._has_value(ws.cell(86, col))
                        ],
                    },
                ],
            }
        ExcelDataExtractor._add_section(output, "risk_prevention_work_details", risk_prevention_work_details)

        critical_assurance: Dict[str, Any] = {}
        critical_assurance_map = {
            "duty_critical": "D90",
            "incident_critical": "D91",
            "availability_assure": "D92",
        }
        for token, addr in critical_assurance_map.items():
            ExcelDataExtractor._put_text(critical_assurance, token, ws[addr])

        posture_rows = []
        for row in range(91, 98):
            raw_category = ws.cell(row, 6)
            raw_attack_count = ws.cell(row, 7)
            raw_defense_rate = ws.cell(row, 8)
            if (
                not ExcelDataExtractor._has_value(raw_category)
                and not ExcelDataExtractor._has_value(raw_attack_count)
                and not ExcelDataExtractor._has_value(raw_defense_rate)
            ):
                continue
            posture_rows.append((raw_category, raw_attack_count, raw_defense_rate))

        if posture_rows:
            critical_assurance["posture_comparison"] = {
                "categories": [ExcelDataExtractor._to_text(row[0]) for row in posture_rows],
                "attack_counts": [ExcelDataExtractor._to_number(row[1], 0) for row in posture_rows],
                "defense_rates": [ExcelDataExtractor._to_number(row[2], 0) for row in posture_rows],
            }
        ExcelDataExtractor._add_section(output, "critical_assurance", critical_assurance)

        output = ExcelDataExtractor._format_numbers_for_output(output)
        output = ExcelDataExtractor._normalize_chart_payload_numbers(output)
        return output

    @staticmethod
    def extract_data(excel_path: Path, template_id: str | None = None) -> Dict[str, Any]:
        """Extract data from Excel file by template-specific extractor."""
        template_config = ExcelDataExtractor._get_template_config(template_id)
        resolved_template_id = template_config["template_id"]
        required_sheet = template_config["required_sheet"]
        extractor = template_config["extractor"]
        try:
            wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
            sheet_names = set(wb.sheetnames)

            if required_sheet not in sheet_names:
                expected_sheets = ", ".join(sorted({cfg["required_sheet"] for cfg in ExcelDataExtractor.TEMPLATE_EXTRACTORS.values()}))
                wb.close()
                raise DataValidationError(
                    field="worksheet",
                    message=(
                        f"Workbook does not match template '{resolved_template_id}'. "
                        f"Expected sheet: {required_sheet}. Available sheets: {', '.join(sorted(sheet_names)) or 'none'}. "
                        f"Registered template sheets: {expected_sheets}."
                    ),
                    template_id=resolved_template_id,
                )

            ws = wb[required_sheet]
            data = extractor(ws)
            wb.close()
            logger.info("Excel parsed with template extractor: %s", resolved_template_id)
            return data
        except InvalidFileException as e:
            raise FileValidationError(
                filename=str(excel_path.name),
                reason=f"Invalid Excel file format: {e}",
            ) from e
        except (DataValidationError, FileValidationError):
            raise
        except Exception as e:
            raise DataValidationError(
                field="excel_data",
                message=f"Excel parsing failed for template '{resolved_template_id}': {e}",
                template_id=resolved_template_id,
            ) from e


ExcelDataExtractor.TEMPLATE_EXTRACTORS = {
    "mss_classic_ops": {
        "template_id": "mss_classic_ops",
        "required_sheet": ExcelDataExtractor.CLASSIC_REQUIRED_SHEET,
        "extractor": ExcelDataExtractor._extract_classic_ops,
    },
    "mss_classic_ops_2": {
        "template_id": "mss_classic_ops_2",
        "required_sheet": ExcelDataExtractor.CLASSIC_REQUIRED_SHEET,
        "extractor": ExcelDataExtractor._extract_classic_ops_2,
    },
}


class ExcelHandler:
    """High-level handler for Excel file uploads."""

    def __init__(self, max_size_mb: int = 10, chunk_size: int = 8192):
        self.validator = ExcelValidator(max_size_mb=max_size_mb)
        self.extractor = ExcelDataExtractor()
        self.chunk_size = chunk_size

    async def process_upload(
        self,
        file_content: bytes,
        filename: str,
        content_type: str,
        session_dir: Path,
        template_id: str | None = None,
    ) -> Dict[str, Any]:
        """Process uploaded Excel file and persist parsed JSON."""
        self.validator.validate_extension(filename)
        self.validator.validate_mime_type(filename, content_type)
        self.validator.validate_size(filename, len(file_content))

        session_dir.mkdir(parents=True, exist_ok=True)
        excel_path = session_dir / "uploaded.xlsx"

        try:
            excel_path.write_bytes(file_content)
            logger.info("File saved: %s (%.2f KB)", excel_path, len(file_content) / 1024)

            if template_id is None:
                logger.warning("Excel upload missing template_id, defaulting to %s", self.extractor.DEFAULT_TEMPLATE_ID)

            input_data = self.extractor.extract_data(excel_path, template_id=template_id)

            json_path = session_dir / "input.json"
            with json_path.open("w", encoding="utf-8") as f:
                json.dump(input_data, f, ensure_ascii=False, indent=2)

            logger.info("JSON saved: %s", json_path)
            return input_data
        except Exception:
            if excel_path.exists():
                try:
                    excel_path.unlink()
                except Exception:
                    pass
            raise

    def get_file_info(self, data: Dict[str, Any], file_size: int) -> Dict[str, Any]:
        """Generate file information summary for API response."""
        preview = {
            "period": data.get("period"),
            "template_id": data.get("template_id", "unknown"),
        }
        if "cover" in data:
            preview["cover"] = data.get("cover")

        return {
            "file_size_mb": round(file_size / 1024 / 1024, 2),
            "preview": preview,
        }
