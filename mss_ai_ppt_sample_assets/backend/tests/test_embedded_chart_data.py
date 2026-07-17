from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

from mss_ai_ppt_sample_assets.backend.modules.ppt_generator import PPTGeneratorV2
from mss_ai_ppt_sample_assets.backend.scripts.embed_linked_chart_data import (
    embed_linked_chart_data,
)


def _add_chart_with_external_workbook():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    chart_data = CategoryChartData()
    chart_data.categories = ["Jan", "Feb"]
    chart_data.add_series("Threats", [10, 20])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(3),
        chart_data,
    ).chart

    embedded_r_id = chart._chartSpace.xlsx_part_rId
    chart.part.rels.pop(embedded_r_id)
    external_r_id = chart.part.relate_to(
        "/Users/example/Desktop/source.xlsx",
        RT.OLE_OBJECT,
        is_external=True,
    )
    chart._chartSpace.externalData.rId = external_r_id
    return presentation, chart


class EmbeddedChartDataTestCase(unittest.TestCase):
    def test_generator_converts_external_workbook_before_replacing_data(self) -> None:
        _, chart = _add_chart_with_external_workbook()
        replacement = CategoryChartData()
        replacement.categories = ["Mar", "Apr"]
        replacement.add_series("Threats", [30, 40])

        generator = PPTGeneratorV2(template_repo=None)
        generator._replace_chart_data(chart, replacement)

        relationship_id = chart._chartSpace.xlsx_part_rId
        relationship = chart.part.rels[relationship_id]
        self.assertFalse(relationship.is_external)
        self.assertEqual(relationship.reltype, RT.PACKAGE)
        self.assertIsNotNone(chart.part.chart_workbook.xlsx_part)
        self.assertEqual(tuple(chart.series[0].values), (30.0, 40.0))

    def test_template_repair_converts_all_external_chart_relationships(self) -> None:
        presentation, _ = _add_chart_with_external_workbook()
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "linked-chart.pptx"
            presentation.save(path)

            converted = embed_linked_chart_data(path)

            self.assertEqual(converted, 1)
            repaired = Presentation(path)
            chart = repaired.slides[0].shapes[0].chart
            relationship_id = chart._chartSpace.xlsx_part_rId
            relationship = chart.part.rels[relationship_id]
            self.assertFalse(relationship.is_external)
            self.assertEqual(relationship.reltype, RT.PACKAGE)
            self.assertIsNotNone(chart.part.chart_workbook.xlsx_part)
            self.assertEqual(tuple(chart.series[0].values), (10.0, 20.0))


if __name__ == "__main__":
    unittest.main()
