from __future__ import annotations

import unittest

from openpyxl import Workbook

from mss_ai_ppt_sample_assets.backend.models.inputs import TenantInput
from mss_ai_ppt_sample_assets.backend.modules.excel_handler import ExcelDataExtractor
from mss_ai_ppt_sample_assets.backend.modules.llm_orchestrator import LLMOrchestratorV2
from mss_ai_ppt_sample_assets.backend.modules.template_loader import TemplateRepository


class PlusGoalReviewPlaceholdersTestCase(unittest.TestCase):
    def test_plus_excel_extractor_outputs_goal_review_sources(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = ExcelDataExtractor.CLASSIC_REQUIRED_SHEET
        ws["J1"] = "示例公司"
        ws["D3"] = "OA系统"
        ws["D4"] = "ERP系统"
        ws["D5"] = "CRM系统"
        ws["G11"] = 12
        ws["G12"] = 34

        extracted = ExcelDataExtractor._extract_classic_ops_2(ws)

        self.assertEqual(extracted["cover"]["company"], "示例公司")
        self.assertEqual(extracted["success_metric"]["core_system_1"], "OA系统")
        self.assertEqual(extracted["success_metric"]["core_system_2"], "ERP系统")
        self.assertEqual(extracted["success_metric"]["core_system_3"], "CRM系统")
        self.assertEqual(extracted["ensure_result"]["risk_total"], "12")
        self.assertEqual(extracted["ensure_result"]["event_total"], "34")

    def test_goal_review_slide_reuses_existing_plus_sources(self) -> None:
        template = TemplateRepository().get_descriptor_v2("mss_classic_ops_2")
        tenant_input = TenantInput(
            raw={
                "cover": {"company": "示例公司"},
                "success_metric": {
                    "core_system_1": "OA系统",
                    "core_system_2": "ERP系统",
                    "core_system_3": "CRM系统",
                },
                "ensure_result": {
                    "risk_total": 12,
                    "event_total": 34,
                },
            }
        )
        orchestrator = LLMOrchestratorV2.__new__(LLMOrchestratorV2)

        placeholders = orchestrator._extract_data_placeholders(tenant_input, template)

        self.assertEqual(
            placeholders["goal_review"],
            {
                "core_system_1": "OA系统",
                "core_system_2": "ERP系统",
                "core_system_3": "CRM系统",
                "company": "示例公司",
            },
        )
        self.assertEqual(placeholders["success_metric"]["company"], "示例公司")
        self.assertEqual(placeholders["ensure_result"]["risk_total"], "12")
        self.assertEqual(placeholders["ensure_result"]["event_total"], "34")
        self.assertEqual(placeholders["key_value_result_coop_update"]["company"], "示例公司")
        self.assertEqual(placeholders["next_phase_action_plan"]["company"], "示例公司")

    def test_shifted_plus_chart_tokens_use_current_page_numbers(self) -> None:
        template = TemplateRepository().get_descriptor_v2("mss_classic_ops_2")
        tenant_input = TenantInput(
            raw={
                "incident_effectiveness": {
                    "response_timeliness": {"labels": ["1h"], "values": [1]},
                    "response_trend": {
                        "months": ["2026-01"],
                        "series": [{"name": "事件数", "values": [1]}],
                    },
                    "incident_distribution": {"categories": ["高危"], "values": [1]},
                },
                "threat_effectiveness": {
                    "attack_source_region_top5": {"categories": ["北京"], "values": [1]},
                    "attack_type_top5": {"categories": ["扫描"], "values": [1]},
                    "externally_attacked_hosts_top5": {"categories": ["host1"], "values": [1]},
                    "threat_trend": {
                        "months": ["2026-01"],
                        "external_attacks": [10000],
                        "malicious_outbound": [20000],
                    },
                },
                "risk_prevention_work_details": {
                    "vulnerability_distribution": {"categories": ["高危"], "values": [1]},
                    "vulnerability_trend": {
                        "months": ["2026-01"],
                        "series": [{"name": "漏洞数", "values": [1]}],
                    },
                },
                "critical_assurance": {
                    "posture_comparison": {
                        "categories": ["国庆"],
                        "attack_counts": [1],
                        "defense_rates": [1],
                    },
                },
            }
        )
        orchestrator = LLMOrchestratorV2.__new__(LLMOrchestratorV2)

        placeholders = orchestrator._extract_data_placeholders(tenant_input, template)

        self.assertEqual(
            {token for token in placeholders["incident_effectiveness"] if token.startswith("P")},
            {"P27_bar", "P27_line", "P27_pie"},
        )
        self.assertEqual(
            {token for token in placeholders["threat_effectiveness"] if token.startswith("P")},
            {"P28_pie_1", "P28_pie_2", "P28_bar", "P28_line"},
        )
        self.assertEqual(
            {token for token in placeholders["risk_prevention_work_details"] if token.startswith("P")},
            {"P29_pie", "P29_line"},
        )
        self.assertEqual(
            {token for token in placeholders["critical_assurance"] if token.startswith("P")},
            {"P30_combo"},
        )


if __name__ == "__main__":
    unittest.main()
