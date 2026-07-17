from __future__ import annotations

import unittest

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from mss_ai_ppt_sample_assets.backend.modules.ppt_generator import (
    PPTGeneratorV2,
    _calculate_nice_axis,
)


class ChartAxisScalingTestCase(unittest.TestCase):
    def test_nice_axis_for_large_mock_values(self) -> None:
        self.assertEqual(_calculate_nice_axis(26680.72), (5000.0, 30000.0))

    def test_nice_axis_handles_small_and_zero_values(self) -> None:
        self.assertEqual(_calculate_nice_axis(1.46), (0.25, 1.5))
        self.assertEqual(_calculate_nice_axis(0), (1.0, 1.0))

    def test_p27_bar_renders_six_intervals_for_mock_data(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        placeholder = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.5), Inches(1), Inches(0.3)
        )
        placeholder.text = "{{P27_bar}}"

        initial_data = CategoryChartData()
        initial_data.categories = ["识别", "响应", "处置", "闭环"]
        initial_data.add_series("时间", [1, 2, 3, 4])
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(1),
            Inches(1),
            Inches(6),
            Inches(3),
            initial_data,
        ).chart

        generator = PPTGeneratorV2(template_repo=None)
        generator._render_p11_bar(
            slide,
            {
                "categories": ["识别", "响应", "处置", "闭环"],
                "series": [
                    {
                        "name": "时间",
                        "values": [1.46, 6.68, 6955.72, 26680.72],
                    }
                ],
            },
            token="P27_bar",
        )

        self.assertEqual(chart.value_axis.minimum_scale, 0.0)
        self.assertEqual(chart.value_axis.maximum_scale, 30000.0)
        self.assertEqual(chart.value_axis.major_unit, 5000.0)
        self.assertEqual(chart.value_axis.tick_labels.number_format, "0")


if __name__ == "__main__":
    unittest.main()
