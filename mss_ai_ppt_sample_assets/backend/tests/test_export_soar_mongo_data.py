from __future__ import annotations

from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from mss_ai_ppt_sample_assets.backend.scripts.export_soar_mongo_data import (
    _write_default_outputs,
    _write_output,
)
from mss_ai_ppt_sample_assets.backend.services.data_ingestion.collectors.mongo_collector import (
    _build_asset_pipeline,
    _transform_asset_doc,
)


class ExportSoarMongoDataTestCase(unittest.TestCase):
    def test_write_output_creates_alarm_and_event_sheets(self) -> None:
        payload = {
            "meta": {"alarm_count": 1, "event_count": 1, "asset_count": 1},
            "alarm": [
                {
                    "alarm_name": "test alarm",
                    "create_time": datetime(2026, 1, 30, 11, 2, 19),
                    "hits": [datetime(2026, 1, 30, 11, 3, 0)],
                }
            ],
            "event": [
                {
                    "event_grading_tag": "重大事件",
                    "识别时长": 1.1,
                    "create_time": datetime(2026, 1, 30, 11, 2, 19),
                }
            ],
            "asset": [
                {
                    "business_name": "核心业务",
                    "asset": "10.0.0.1",
                    "level": "L3",
                    "is_service": 1,
                    "asset_type": "server",
                    "security_domain": "dmz",
                }
            ],
        }

        with TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "soar_export.xlsx"
            _write_output(payload, output_path)

            workbook = load_workbook(output_path)
            self.assertEqual(workbook.sheetnames, ["告警表", "事件表", "资产表"])
            self.assertEqual(workbook["告警表"]["A1"].value, "alarm_name")
            self.assertEqual(workbook["告警表"]["A2"].value, "test alarm")
            self.assertEqual(workbook["告警表"]["C2"].value, '["2026-01-30T11:03:00"]')
            self.assertEqual(workbook["事件表"]["A1"].value, "event_grading_tag")
            self.assertEqual(workbook["事件表"]["B1"].value, "识别时长")
            self.assertEqual(workbook["事件表"]["B2"].value, 1.1)
            self.assertEqual(workbook["资产表"]["A1"].value, "business_name")
            self.assertEqual(workbook["资产表"]["A2"].value, "核心业务")
            self.assertEqual(workbook["资产表"]["F2"].value, "dmz")

    def test_write_default_outputs_creates_json_and_xlsx(self) -> None:
        payload = {"meta": {}, "alarm": [], "event": [], "asset": []}

        with TemporaryDirectory() as tmpdir:
            json_output_path, xlsx_output_path = _write_default_outputs(payload, Path(tmpdir))

            self.assertTrue(json_output_path.exists())
            self.assertTrue(xlsx_output_path.exists())
            self.assertEqual(json_output_path.name, "soar_raw_export.json")
            self.assertEqual(xlsx_output_path.name, "soar_raw_export.xlsx")

    def test_build_asset_pipeline_matches_expected_shape(self) -> None:
        pipeline = _build_asset_pipeline("tenant-1", "Business")

        self.assertEqual(
            pipeline,
            [
                {"$match": {"is_deleted": 0, "company_id": "tenant-1"}},
                {
                    "$lookup": {
                        "from": "Business",
                        "localField": "business_id",
                        "foreignField": "_id",
                        "as": "business",
                    }
                },
                {"$unwind": {"path": "$business", "preserveNullAndEmptyArrays": True}},
                {
                    "$project": {
                        "_id": 0,
                        "business_name": {"$ifNull": ["$business.business_name", ""]},
                        "asset": "$asset",
                        "level": {"$ifNull": ["$business.level", ""]},
                        "is_service": "$is_service",
                        "asset_type": "$asset_type",
                        "security_domain": {"$ifNull": ["$security_domain", ""]},
                    }
                },
            ],
        )

    def test_transform_asset_doc_keeps_only_output_fields(self) -> None:
        transformed = _transform_asset_doc(
            {
                "business_name": "业务A",
                "asset": "app01",
                "level": "L2",
                "is_service": 0,
                "asset_type": "vm",
                "security_domain": "prod",
                "ignored": "value",
            }
        )

        self.assertEqual(
            transformed,
            {
                "business_name": "业务A",
                "asset": "app01",
                "level": "L2",
                "is_service": 0,
                "asset_type": "vm",
                "security_domain": "prod",
            },
        )


if __name__ == "__main__":
    unittest.main()
