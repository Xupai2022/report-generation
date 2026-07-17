from __future__ import annotations

import argparse
import io
import os
import tempfile
import zipfile
from copy import copy
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPE_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
PACKAGE_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"
)
XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

NS = {"c": CHART_NS, "r": DOC_REL_NS}
ET.register_namespace("c", CHART_NS)
ET.register_namespace("r", DOC_REL_NS)


def _cached_values(reference: ET.Element | None) -> tuple[list[object], str | None]:
    if reference is None:
        return [], None

    cache = reference.find("c:strCache", NS)
    if cache is None:
        cache = reference.find("c:numCache", NS)
    if cache is None:
        return [], None

    format_code = cache.findtext("c:formatCode", default=None, namespaces=NS)
    points: dict[int, object] = {}
    for point in cache.findall("c:pt", NS):
        index = int(point.get("idx", "0"))
        value = point.findtext("c:v", default="", namespaces=NS)
        if reference.tag == f"{{{CHART_NS}}}numRef" and value != "":
            try:
                value = float(value)
            except ValueError:
                pass
        points[index] = value

    point_count_text = cache.findtext("c:ptCount", default="0", namespaces=NS)
    point_count = int(point_count_text or 0)
    if points:
        point_count = max(point_count, max(points) + 1)
    return [points.get(index) for index in range(point_count)], format_code


def _series_name(series: ET.Element, fallback: str) -> str:
    name = series.findtext("c:tx/c:strRef/c:strCache/c:pt/c:v", default=None, namespaces=NS)
    if name is None:
        name = series.findtext("c:tx/c:v", default=None, namespaces=NS)
    return fallback if name is None else name


def _build_embedded_workbook(chart_root: ET.Element) -> bytes:
    series_elements = chart_root.findall(".//c:ser", NS)
    if not series_elements:
        raise ValueError("chart has no cached series data")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"

    first_category_values: list[object] = []
    first_category_format: str | None = None
    max_point_count = 0

    for series_index, series in enumerate(series_elements, start=2):
        column = get_column_letter(series_index)
        worksheet.cell(1, series_index, _series_name(series, f"Series {series_index - 1}"))

        category_reference = series.find("c:cat/c:strRef", NS)
        if category_reference is None:
            category_reference = series.find("c:cat/c:numRef", NS)
        category_values, category_format = _cached_values(category_reference)
        if not first_category_values and category_values:
            first_category_values = category_values
            first_category_format = category_format

        value_reference = series.find("c:val/c:numRef", NS)
        values, value_format = _cached_values(value_reference)
        max_point_count = max(max_point_count, len(category_values), len(values))

        for row_index, value in enumerate(values, start=2):
            cell = worksheet.cell(row_index, series_index, value)
            if value_format and value_format != "General":
                cell.number_format = value_format

        name_formula = series.find("c:tx/c:strRef/c:f", NS)
        if name_formula is not None:
            name_formula.text = f"Sheet1!${column}$1"

        category_formula = series.find("c:cat/c:strRef/c:f", NS)
        if category_formula is None:
            category_formula = series.find("c:cat/c:numRef/c:f", NS)
        if category_formula is not None:
            category_formula.text = f"Sheet1!$A$2:$A${max(len(category_values), len(values)) + 1}"

        value_formula = series.find("c:val/c:numRef/c:f", NS)
        if value_formula is not None:
            value_formula.text = f"Sheet1!${column}$2:${column}${len(values) + 1}"

    for row_index in range(max_point_count):
        value = first_category_values[row_index] if row_index < len(first_category_values) else None
        cell = worksheet.cell(row_index + 2, 1, value)
        if first_category_format and first_category_format != "General":
            cell.number_format = first_category_format

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _chart_relationship_path(chart_path: str) -> str:
    chart_name = chart_path.rsplit("/", 1)[-1]
    return f"ppt/charts/_rels/{chart_name}.rels"


def _embed_chart(
    chart_path: str,
    chart_xml: bytes,
    relationship_xml: bytes,
    embedding_index: int,
) -> tuple[bytes, bytes, str, bytes] | None:
    chart_root = ET.fromstring(chart_xml)
    external_data = chart_root.find("c:externalData", NS)
    if external_data is None:
        return None

    relationship_id = external_data.get(f"{{{DOC_REL_NS}}}id")
    if not relationship_id:
        return None

    relationship_root = ET.fromstring(relationship_xml)
    relationship = next(
        (
            item
            for item in relationship_root.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
            if item.get("Id") == relationship_id
        ),
        None,
    )
    if relationship is None or relationship.get("TargetMode") != "External":
        return None

    workbook_blob = _build_embedded_workbook(chart_root)
    embedding_name = f"Microsoft_Excel_Worksheet{embedding_index}.xlsx"
    embedding_path = f"ppt/embeddings/{embedding_name}"

    relationship.set("Type", PACKAGE_REL_TYPE)
    relationship.set("Target", f"../embeddings/{embedding_name}")
    relationship.attrib.pop("TargetMode", None)

    updated_chart_xml = ET.tostring(chart_root, encoding="utf-8", xml_declaration=True)
    updated_relationship_xml = ET.tostring(
        relationship_root,
        encoding="utf-8",
        xml_declaration=True,
    )
    return updated_chart_xml, updated_relationship_xml, embedding_path, workbook_blob


def _ensure_xlsx_content_type(content_types_xml: bytes) -> bytes:
    root = ET.fromstring(content_types_xml)
    defaults = root.findall(f"{{{CONTENT_TYPE_NS}}}Default")
    if any(item.get("Extension", "").lower() == "xlsx" for item in defaults):
        return content_types_xml

    ET.SubElement(
        root,
        f"{{{CONTENT_TYPE_NS}}}Default",
        {"Extension": "xlsx", "ContentType": XLSX_CONTENT_TYPE},
    )
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def embed_linked_chart_data(pptx_path: Path) -> int:
    """Convert every externally linked chart workbook in *pptx_path* to embedded data."""
    pptx_path = pptx_path.resolve()
    with zipfile.ZipFile(pptx_path, "r") as source:
        entries = {item.filename: source.read(item.filename) for item in source.infolist()}
        entry_info = {item.filename: copy(item) for item in source.infolist()}

    replacements: dict[str, bytes] = {}
    embedding_index = 1
    converted = 0
    chart_paths = sorted(
        name
        for name in entries
        if name.startswith("ppt/charts/chart") and name.endswith(".xml")
    )

    for chart_path in chart_paths:
        relationship_path = _chart_relationship_path(chart_path)
        if relationship_path not in entries:
            continue
        result = _embed_chart(
            chart_path,
            entries[chart_path],
            entries[relationship_path],
            embedding_index,
        )
        if result is None:
            continue

        chart_xml, relationship_xml, embedding_path, workbook_blob = result
        replacements[chart_path] = chart_xml
        replacements[relationship_path] = relationship_xml
        replacements[embedding_path] = workbook_blob
        embedding_index += 1
        converted += 1

    if converted == 0:
        return 0

    replacements["[Content_Types].xml"] = _ensure_xlsx_content_type(
        entries["[Content_Types].xml"]
    )

    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{pptx_path.stem}-embed-",
        suffix=".pptx",
        dir=pptx_path.parent,
    )
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)

    try:
        with zipfile.ZipFile(temporary_path, "w") as target:
            for name, data in entries.items():
                target.writestr(entry_info[name], replacements.get(name, data))
            for name, data in replacements.items():
                if name not in entries:
                    target.writestr(name, data, compress_type=zipfile.ZIP_DEFLATED)

        with zipfile.ZipFile(temporary_path, "r") as check:
            bad_entry = check.testzip()
            if bad_entry:
                raise ValueError(f"corrupt ZIP entry after conversion: {bad_entry}")
        os.replace(temporary_path, pptx_path)
    finally:
        temporary_path.unlink(missing_ok=True)

    return converted


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Convert externally linked PowerPoint chart workbooks to embedded XLSX data."
    )
    parser.add_argument("pptx", nargs="+", type=Path)
    arguments = parser.parse_args()

    for pptx_path in arguments.pptx:
        converted = embed_linked_chart_data(pptx_path)
        print(f"{pptx_path}: converted {converted} linked chart(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
